import io
import logging
import re
import uuid
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor
from fastapi import HTTPException, status
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfgen import canvas
from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.approval import Approval
from app.models.report_version import ReportVersion
from app.models.review_cycle import ReviewCycle
from app.repositories.report_governance_repository import ReportGovernanceRepository
from app.repositories.review_cycle_repository import ReviewCycleRepository

logger = logging.getLogger(__name__)


class NumberedCanvas(canvas.Canvas):
    """
    Two-pass canvas that automatically computes and stamps 'Page X of Y'
    along with running headers and footers on every page.
    """

    def __init__(self, *args: Any, **kwargs: Any):
        super().__init__(*args, **kwargs)
        self._saved_page_states: List[Dict[str, Any]] = []

    def showPage(self) -> None:
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self) -> None:
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count: int) -> None:
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))

        # Running Header
        self.drawString(54, 752, "PERFORMai-360 — AI-Driven Bias-Aware 360° Evaluation")
        self.setStrokeColor(colors.HexColor("#E2E8F0"))
        self.setLineWidth(0.75)
        self.line(54, 744, 558, 744)

        # Running Footer
        self.line(54, 45, 558, 45)
        self.drawString(54, 32, "CONFIDENTIAL — FOR AUTHORIZED PERFORMANCE REVIEW ONLY")
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 32, page_str)
        self.restoreState()


class ReportExportService:
    """
    Production-grade PDF and DOCX export generator for Phase 9.
    Renders immutable report snapshots with enterprise styling,
    exact version identification, human sign-off records, and citation grounding.
    """

    def __init__(self, db: AsyncSession):
        self.db = db

    async def _resolve_report_and_version(
        self,
        report_id: uuid.UUID,
        version_number: Optional[int] = None,
        version_id: Optional[uuid.UUID] = None,
        viewer_id: Optional[uuid.UUID] = None,
        viewer_role: Optional[str] = None,
    ) -> Dict[str, Any]:
        report = await ReportGovernanceRepository.get_report_by_id(self.db, report_id)
        if not report:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Report '{report_id}' not found.",
            )

        cycle = report.review_cycle
        if not cycle:
            cycle = await ReviewCycleRepository.get_by_id(self.db, report.cycle_id)

        if not cycle:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Associated review cycle not found.",
            )

        # Scoping verification
        role_lower = (viewer_role or "employee").lower()
        if role_lower == "employee":
            if not viewer_id or cycle.employee_id != viewer_id:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail=f"Report '{report_id}' not found.",
                )

        # Resolve exact version
        target_version: Optional[ReportVersion] = None
        if version_number is not None:
            if version_number <= 0:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
                    detail="version_number must be greater than 0.",
                )
            target_version = await ReportGovernanceRepository.get_version_by_number(self.db, report_id, version_number)
            if not target_version:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail=f"Version '{version_number}' not found for report '{report_id}'.",
                )
        elif version_id is not None:
            target_version = await ReportGovernanceRepository.get_version_by_id(self.db, version_id)
            if not target_version or target_version.report_id != report_id:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail=f"Version ID '{version_id}' not found for report '{report_id}'.",
                )
        else:
            # Latest snapshot
            target_version = await ReportGovernanceRepository.get_latest_version(self.db, report_id)

        content = target_version.structured_content if target_version else (report.structured_content or {})
        ver_num = target_version.version_number if target_version else report.version_number
        ver_status = target_version.status if target_version else report.status
        is_ai = target_version.is_ai_generated if target_version else False

        # Look up approval sign-off record for this version
        approvals = await ReportGovernanceRepository.list_approvals_for_report(self.db, report_id)
        matching_approval: Optional[Approval] = None
        for a in approvals:
            if a.version_number == ver_num and a.decision == "approved":
                matching_approval = a
                break

        # If requester is employee, mask confidential reviewer notes
        safe_content = dict(content or {})
        if role_lower == "employee":
            safe_content.pop("reviewer_notes", None)

        return {
            "report": report,
            "cycle": cycle,
            "version_number": ver_num,
            "status": ver_status,
            "is_ai_generated": is_ai,
            "content": safe_content,
            "approval": matching_approval,
            "viewer_role": role_lower,
        }

    # ---------------------------------------------------------------------------
    # PDF Generation
    # ---------------------------------------------------------------------------

    async def generate_pdf(
        self,
        report_id: uuid.UUID,
        version_number: Optional[int] = None,
        version_id: Optional[uuid.UUID] = None,
        viewer_id: Optional[uuid.UUID] = None,
        viewer_role: Optional[str] = None,
    ) -> Tuple[bytes, str]:
        data = await self._resolve_report_and_version(
            report_id=report_id,
            version_number=version_number,
            version_id=version_id,
            viewer_id=viewer_id,
            viewer_role=viewer_role,
        )

        cycle: ReviewCycle = data["cycle"]
        content: Dict[str, Any] = data["content"]
        ver_num: int = data["version_number"]
        rep_status: str = data["status"]
        approval: Optional[Approval] = data["approval"]

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=54,
            rightMargin=54,
            topMargin=72,
            bottomMargin=54,
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=20,
            leading=24,
            textColor=colors.HexColor("#0F172A"),
        )
        subtitle_style = ParagraphStyle(
            "DocSubTitle",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=11,
            leading=15,
            textColor=colors.HexColor("#64748B"),
        )
        h1_style = ParagraphStyle(
            "Heading1Custom",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=13,
            leading=17,
            textColor=colors.HexColor("#1E293B"),
            spaceBefore=14,
            spaceAfter=6,
        )
        body_style = ParagraphStyle(
            "BodyCustom",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=14,
            textColor=colors.HexColor("#334155"),
        )
        meta_label = ParagraphStyle(
            "MetaLabel",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#475569"),
        )
        meta_val = ParagraphStyle(
            "MetaVal",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor("#0F172A"),
        )

        story: List[Any] = []

        # 1. Header Title
        story.append(Paragraph("PERFORMai-360 Evaluation Report", title_style))
        story.append(Spacer(1, 4))
        story.append(
            Paragraph("AI-Assisted, Evidence-Grounded Performance Review with Immutable Human Sign-Off", subtitle_style)
        )
        story.append(Spacer(1, 12))
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#4F46E5"), spaceAfter=14))

        # 2. Metadata Grid Table
        emp_name = cycle.employee.full_name if cycle.employee else "Unknown Employee"
        emp_email = cycle.employee.email if cycle.employee else "—"
        mgr_name = cycle.owner.full_name if cycle.owner else "Management"
        period_str = f"{cycle.period_start} to {cycle.period_end}"

        meta_data = [
            [
                Paragraph("Employee:", meta_label),
                Paragraph(f"{emp_name} ({emp_email})", meta_val),
                Paragraph("Review Lead:", meta_label),
                Paragraph(mgr_name, meta_val),
            ],
            [
                Paragraph("Review Cycle:", meta_label),
                Paragraph(cycle.title, meta_val),
                Paragraph("Report Version:", meta_label),
                Paragraph(f"Version {ver_num}", meta_val),
            ],
            [
                Paragraph("Review Period:", meta_label),
                Paragraph(period_str, meta_val),
                Paragraph("Governance Status:", meta_label),
                Paragraph(rep_status.upper().replace("_", " "), meta_val),
            ],
        ]
        meta_table = Table(meta_data, colWidths=[90, 190, 90, 134])
        meta_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                    ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                    ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ]
            )
        )
        story.append(meta_table)
        story.append(Spacer(1, 12))

        # 3. Governance Status / Approval Callout Box
        if rep_status == "approved" and approval:
            app_text = (
                f"<b>OFFICIAL HUMAN APPROVAL RECORD:</b> Approved and signed off by <b>{approval.reviewer.full_name if approval.reviewer else 'Authorized Manager'}</b> "
                f"({approval.reviewer.role.value if approval.reviewer and hasattr(approval.reviewer.role, 'value') else 'Manager'}) on "
                f"{approval.created_at.strftime('%Y-%m-%d at %H:%M UTC')}.<br/>"
                f'<i>Sign-off Comments:</i> "{approval.comments or "Report vetted and approved for delivery."}"'
            )
            banner_bg = colors.HexColor("#ECFDF5")
            banner_border = colors.HexColor("#10B981")
            banner_txt_color = colors.HexColor("#065F46")
        else:
            app_text = (
                "<b>GOVERNANCE STATUS: DRAFT (UNAPPROVED)</b><br/>"
                "This performance evaluation document is an active working draft. It has NOT been finalized by an authorized "
                "manager and cannot be used as an official employment record until formal human sign-off."
            )
            banner_bg = colors.HexColor("#FFFBEB")
            banner_border = colors.HexColor("#F59E0B")
            banner_txt_color = colors.HexColor("#92400E")

        banner_p = Paragraph(
            app_text,
            ParagraphStyle("BannerText", parent=body_style, textColor=banner_txt_color, fontSize=9, leading=13),
        )
        banner_table = Table([[banner_p]], colWidths=[504])
        banner_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), banner_bg),
                    ("BOX", (0, 0), (-1, -1), 1, banner_border),
                    ("TOPPADDING", (0, 0), (-1, -1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                    ("LEFTPADDING", (0, 0), (-1, -1), 10),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ]
            )
        )
        story.append(banner_table)
        story.append(Spacer(1, 14))

        # 4. Executive Summary
        story.append(Paragraph("1. Executive Summary", h1_style))
        exec_summary = content.get("executive_summary") or "No executive summary supplied."
        story.append(Paragraph(exec_summary, body_style))
        story.append(Spacer(1, 10))

        # 5. Validated Strengths
        strengths = content.get("strengths", [])
        story.append(Paragraph(f"2. Validated Strengths ({len(strengths)})", h1_style))
        if strengths:
            for s in strengths:
                theme = s.get("theme") or s.get("title") or "Core Strength"
                narrative = s.get("narrative") or s.get("description") or ""
                evidence_ids = s.get("supporting_evidence_ids", [])
                conf = s.get("confidence_score")
                conf_str = f" · Confidence: {int(conf * 100)}%" if conf is not None else ""

                item_text = f"<b>• {theme}</b>{conf_str}<br/>{narrative}"
                if evidence_ids:
                    cites = ", ".join([str(cid)[:8] + "..." for cid in evidence_ids])
                    item_text += f"<br/><font color='#64748B' size='8'>Supporting Citations: {cites}</font>"

                story.append(Paragraph(item_text, body_style))
                story.append(Spacer(1, 6))
        else:
            story.append(Paragraph("No validated strengths recorded.", body_style))
        story.append(Spacer(1, 10))

        # 6. Growth Opportunities
        growth = content.get("growth_areas", [])
        story.append(Paragraph(f"3. Growth & Development Opportunities ({len(growth)})", h1_style))
        if growth:
            for g in growth:
                theme = g.get("theme") or g.get("title") or "Growth Area"
                narrative = g.get("narrative") or g.get("development_suggestion") or g.get("gap_observed") or ""
                item_text = f"<b>• {theme}</b><br/>{narrative}"
                story.append(Paragraph(item_text, body_style))
                story.append(Spacer(1, 6))
        else:
            story.append(Paragraph("No growth areas recorded.", body_style))
        story.append(Spacer(1, 10))

        # 7. Goal Progress
        goals = content.get("goal_progress", [])
        if goals:
            story.append(Paragraph("4. Goal Progress & Outcomes", h1_style))
            for goal in goals:
                desc = goal.get("description", "Goal")
                stat = goal.get("status", "in_progress").upper()
                target = goal.get("measurable_target")
                target_str = f" (Target: {target})" if target else ""
                notes = goal.get("outcome_notes", "")
                g_text = f"<b>• [{stat}] {desc}</b>{target_str}"
                if notes:
                    g_text += f"<br/>{notes}"
                story.append(Paragraph(g_text, body_style))
                story.append(Spacer(1, 4))
            story.append(Spacer(1, 10))

        # 8. Impact Highlights
        impacts = content.get("impact_highlights", [])
        if impacts:
            story.append(Paragraph("5. Impact Highlights", h1_style))
            for imp in impacts:
                title = imp.get("title", "Impact Item")
                summary = imp.get("impact_summary", "")
                story.append(Paragraph(f"<b>• {title}:</b> {summary}", body_style))
                story.append(Spacer(1, 4))
            story.append(Spacer(1, 10))

        # 9. Evidence Limitations & Discussion Questions
        gaps = content.get("evidence_gaps_and_limitations", [])
        if gaps:
            story.append(Paragraph("6. Evidence Gaps & Limitations", h1_style))
            for gap in gaps:
                story.append(Paragraph(f"• {gap}", body_style))
                story.append(Spacer(1, 3))
            story.append(Spacer(1, 10))

        questions = content.get("suggested_discussion_questions", [])
        if questions:
            story.append(Paragraph("7. Discussion Questions for Review Check-In", h1_style))
            for q in questions:
                q_text = q if isinstance(q, str) else q.get("question", str(q))
                story.append(Paragraph(f"• {q_text}", body_style))
                story.append(Spacer(1, 3))
            story.append(Spacer(1, 10))

        # 10. Reviewer Notes (Excluded if employee)
        reviewer_notes = content.get("reviewer_notes")
        if reviewer_notes and data["viewer_role"] != "employee":
            story.append(
                KeepTogether(
                    [
                        Paragraph("8. Reviewer & Calibration Notes (Confidential)", h1_style),
                        Paragraph(f"<i>{reviewer_notes}</i>", body_style),
                    ]
                )
            )

        # Build Document
        doc.build(story, canvasmaker=NumberedCanvas)
        pdf_bytes = buffer.getvalue()
        buffer.close()

        clean_emp = re.sub(r"[^a-zA-Z0-9_-]", "_", emp_name.replace(" ", "_"))
        filename = f"PERFORMai360_Report_{clean_emp}_v{ver_num}.pdf"

        return pdf_bytes, filename

    # ---------------------------------------------------------------------------
    # DOCX Generation
    # ---------------------------------------------------------------------------

    async def generate_docx(
        self,
        report_id: uuid.UUID,
        version_number: Optional[int] = None,
        version_id: Optional[uuid.UUID] = None,
        viewer_id: Optional[uuid.UUID] = None,
        viewer_role: Optional[str] = None,
    ) -> Tuple[bytes, str]:
        data = await self._resolve_report_and_version(
            report_id=report_id,
            version_number=version_number,
            version_id=version_id,
            viewer_id=viewer_id,
            viewer_role=viewer_role,
        )

        cycle: ReviewCycle = data["cycle"]
        content: Dict[str, Any] = data["content"]
        ver_num: int = data["version_number"]
        rep_status: str = data["status"]
        approval: Optional[Approval] = data["approval"]

        emp_name = cycle.employee.full_name if cycle.employee else "Unknown Employee"
        emp_email = cycle.employee.email if cycle.employee else "—"
        mgr_name = cycle.owner.full_name if cycle.owner else "Management"
        period_str = f"{cycle.period_start} to {cycle.period_end}"

        doc = Document()

        # Set 1-inch margins
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # Document Title
        p_title = doc.add_paragraph()
        run_title = p_title.add_run("PERFORMai-360 Evaluation Report")
        run_title.font.name = "Arial"
        run_title.font.size = Pt(20)
        run_title.font.bold = True
        run_title.font.color.rgb = RGBColor(15, 23, 42)
        p_title.paragraph_format.space_after = Pt(2)

        # Subtitle
        p_sub = doc.add_paragraph()
        run_sub = p_sub.add_run("AI-Assisted, Evidence-Grounded Performance Review with Immutable Human Sign-Off")
        run_sub.font.name = "Arial"
        run_sub.font.size = Pt(10)
        run_sub.font.color.rgb = RGBColor(100, 116, 139)
        p_sub.paragraph_format.space_after = Pt(12)

        # Metadata Table
        table = doc.add_table(rows=3, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        meta_rows = [
            ("Employee:", f"{emp_name} ({emp_email})", "Review Lead:", mgr_name),
            ("Review Cycle:", cycle.title, "Report Version:", f"Version {ver_num}"),
            ("Review Period:", period_str, "Governance Status:", rep_status.upper().replace("_", " ")),
        ]

        for r_idx, row in enumerate(meta_rows):
            for c_idx, val in enumerate(row):
                cell = table.cell(r_idx, c_idx)
                cell_p = cell.paragraphs[0]
                cell_run = cell_p.add_run(val)
                cell_run.font.name = "Arial"
                cell_run.font.size = Pt(8.5)
                if c_idx % 2 == 0:
                    cell_run.font.bold = True
                    cell_run.font.color.rgb = RGBColor(71, 85, 105)
                else:
                    cell_run.font.color.rgb = RGBColor(15, 23, 42)
                # Subtle cell shading
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
                cell._tc.get_or_add_tcPr().append(shd)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

        # Approval Banner Box
        banner_p = doc.add_paragraph()
        if rep_status == "approved" and approval:
            b_run = banner_p.add_run(
                f"OFFICIAL HUMAN APPROVAL RECORD: Approved by "
                f"{approval.reviewer.full_name if approval.reviewer else 'Authorized Manager'} on "
                f"{approval.created_at.strftime('%Y-%m-%d %H:%M UTC')}. "
                f'Comments: "{approval.comments or "Report vetted and approved."}"'
            )
            b_run.font.color.rgb = RGBColor(6, 95, 70)
        else:
            b_run = banner_p.add_run(
                "GOVERNANCE STATUS: DRAFT (UNAPPROVED) — This document is an active working draft pending official manager sign-off."
            )
            b_run.font.color.rgb = RGBColor(146, 64, 14)
        b_run.font.name = "Arial"
        b_run.font.bold = True
        b_run.font.size = Pt(9)
        banner_p.paragraph_format.space_after = Pt(14)

        def add_heading(text: str) -> None:
            h = doc.add_paragraph()
            hr = h.add_run(text)
            hr.font.name = "Arial"
            hr.font.size = Pt(13)
            hr.font.bold = True
            hr.font.color.rgb = RGBColor(30, 41, 59)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)

        def add_body(text: str) -> None:
            p = doc.add_paragraph()
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(51, 65, 85)
            p.paragraph_format.space_after = Pt(6)

        # 1. Executive Summary
        add_heading("1. Executive Summary")
        add_body(content.get("executive_summary") or "No executive summary supplied.")

        # 2. Validated Strengths
        strengths = content.get("strengths", [])
        add_heading(f"2. Validated Strengths ({len(strengths)})")
        for s in strengths:
            theme = s.get("theme") or s.get("title") or "Core Strength"
            narrative = s.get("narrative") or s.get("description") or ""
            evidence_ids = s.get("supporting_evidence_ids", [])
            conf = s.get("confidence_score")
            conf_str = f" [Confidence: {int(conf * 100)}%]" if conf is not None else ""

            p = doc.add_paragraph(style="List Bullet")
            r1 = p.add_run(f"{theme}{conf_str}: ")
            r1.font.bold = True
            r1.font.name = "Arial"
            r1.font.size = Pt(9.5)
            r2 = p.add_run(narrative)
            r2.font.name = "Arial"
            r2.font.size = Pt(9.5)

            if evidence_ids:
                p_cites = doc.add_paragraph()
                p_cites.paragraph_format.left_indent = Inches(0.25)
                c_run = p_cites.add_run(f"Citations: {', '.join([str(c)[:8] + '...' for c in evidence_ids])}")
                c_run.font.name = "Arial"
                c_run.font.size = Pt(8)
                c_run.font.color.rgb = RGBColor(100, 116, 139)

        # 3. Growth Areas
        growth = content.get("growth_areas", [])
        add_heading(f"3. Growth & Development Opportunities ({len(growth)})")
        for g in growth:
            theme = g.get("theme") or g.get("title") or "Opportunity"
            narrative = g.get("narrative") or g.get("development_suggestion") or g.get("gap_observed") or ""
            p = doc.add_paragraph(style="List Bullet")
            r1 = p.add_run(f"{theme}: ")
            r1.font.bold = True
            r1.font.name = "Arial"
            r1.font.size = Pt(9.5)
            r2 = p.add_run(narrative)
            r2.font.name = "Arial"
            r2.font.size = Pt(9.5)

        # 4. Goals Progress
        goals = content.get("goal_progress", [])
        if goals:
            add_heading("4. Goal Progress & Outcomes")
            for goal in goals:
                desc = goal.get("description", "Goal")
                stat = goal.get("status", "in_progress").upper()
                p = doc.add_paragraph(style="List Bullet")
                r = p.add_run(f"[{stat}] {desc}")
                r.font.name = "Arial"
                r.font.size = Pt(9.5)

        # 5. Reviewer Notes (Manager only)
        reviewer_notes = content.get("reviewer_notes")
        if reviewer_notes and data["viewer_role"] != "employee":
            add_heading("5. Reviewer & Calibration Notes (Confidential)")
            add_body(reviewer_notes)

        # Output to buffer
        out_buffer = io.BytesIO()
        doc.save(out_buffer)
        docx_bytes = out_buffer.getvalue()
        out_buffer.close()

        clean_emp = re.sub(r"[^a-zA-Z0-9_-]", "_", emp_name.replace(" ", "_"))
        filename = f"PERFORMai360_Report_{clean_emp}_v{ver_num}.docx"

        return docx_bytes, filename
