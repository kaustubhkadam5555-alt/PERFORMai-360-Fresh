from typing import BinaryIO

import docx
from docx.opc.exceptions import PackageNotFoundError

from app.parsers.base import DocumentParser, ParsedDocument

MAX_DOCX_PARAGRAPHS = 2_000
MAX_DOCX_CHARS = 50_000


class DocxParser(DocumentParser):
    def parse(self, stream: BinaryIO, filename: str) -> ParsedDocument:
        try:
            doc = docx.Document(stream)
        except (PackageNotFoundError, Exception) as exc:
            raise ValueError(f"Corrupted or invalid DOCX file: {exc}") from exc

        paragraphs_text = []
        total_chars = 0
        truncated = False
        paragraph_count = 0

        # Extract text from standard paragraphs
        for p in doc.paragraphs:
            if paragraph_count >= MAX_DOCX_PARAGRAPHS:
                truncated = True
                break
            text = p.text.strip()
            if text:
                paragraph_count += 1
                paragraphs_text.append(text)
                total_chars += len(text)
                if total_chars >= MAX_DOCX_CHARS:
                    truncated = True
                    break

        # Extract text from tables if within character limit
        if total_chars < MAX_DOCX_CHARS:
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        row_line = " | ".join(row_cells)
                        paragraphs_text.append(row_line)
                        total_chars += len(row_line)
                        if total_chars >= MAX_DOCX_CHARS:
                            truncated = True
                            break
                if truncated:
                    break

        full_text = "\n\n".join(paragraphs_text)
        if len(full_text) > MAX_DOCX_CHARS:
            full_text = full_text[:MAX_DOCX_CHARS]
            truncated = True

        words = len(full_text.split())
        snippet = full_text[:500].strip()

        return ParsedDocument(
            text=full_text,
            snippet=snippet,
            word_count=words,
            char_count=len(full_text),
            page_count=None,
            metadata={
                "format": "docx",
                "filename": filename,
                "paragraph_count": paragraph_count,
                "table_count": len(doc.tables),
                "is_truncated": truncated,
            },
        )
