import io
import os
import uuid
from unittest.mock import patch

import docx
import pytest
from httpx import AsyncClient
from pypdf import PdfWriter

from app.core.config import settings


def generate_pdf_bytes(title: str = "Quarterly Evaluation") -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def generate_docx_bytes(title: str = "Design RFC", content: str = "Architecture details") -> bytes:
    doc = docx.Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(content)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
async def setup_upload_env(async_client: AsyncClient):
    emp_res = await async_client.post(
        "/api/v1/users",
        json={"email": f"eng_{uuid.uuid4().hex[:8]}@example.com", "full_name": "Engineer Eve", "role": "employee"},
    )
    employee_id = emp_res.json()["id"]

    mgr_res = await async_client.post(
        "/api/v1/users",
        json={"email": f"mgr_{uuid.uuid4().hex[:8]}@example.com", "full_name": "Manager Bob", "role": "manager"},
    )
    owner_id = mgr_res.json()["id"]

    cycle_res = await async_client.post(
        "/api/v1/review-cycles",
        json={
            "title": "Eve 2026 Evaluation",
            "employee_id": employee_id,
            "owner_id": owner_id,
            "period_start": "2026-01-01",
            "period_end": "2026-12-31",
        },
    )
    cycle_id = cycle_res.json()["id"]

    return {
        "employee_id": employee_id,
        "owner_id": owner_id,
        "cycle_id": cycle_id,
    }


@pytest.mark.asyncio
async def test_upload_text_evidence(async_client: AsyncClient, setup_upload_env):
    env = setup_upload_env
    cycle_id = env["cycle_id"]
    employee_id = env["employee_id"]

    file_content = b"Completed high-priority performance optimization on search index.\nLatency dropped by 45%."
    files = {"file": ("optimization_notes.txt", file_content, "text/plain")}
    data = {
        "cycle_id": cycle_id,
        "employee_id": employee_id,
        "title": "Search Optimization Notes",
        "source_type": "note",
        "access_policy": "internal",
    }

    res = await async_client.post("/api/v1/evidence/upload", files=files, data=data)
    assert res.status_code == 201
    ev = res.json()
    assert ev["title"] == "Search Optimization Notes"
    assert ev["source_type"] == "note"
    assert ev["access_policy"] == "internal"
    assert ev["storage_reference"] is not None
    assert os.path.exists(ev["storage_reference"])

    metadata = ev["metadata_json"]
    assert metadata["original_filename"] == "optimization_notes.txt"
    assert metadata["detected_format"] == "txt"
    assert metadata["word_count"] > 5
    assert "Search index" in metadata["snippet"] or "optimization" in metadata["snippet"].lower()


@pytest.mark.asyncio
async def test_upload_pdf_evidence_via_cycle_endpoint(async_client: AsyncClient, setup_upload_env):
    env = setup_upload_env
    cycle_id = env["cycle_id"]

    pdf_bytes = generate_pdf_bytes("System Architecture RFC")
    files = {"file": ("architecture_rfc.pdf", pdf_bytes, "application/pdf")}
    data = {
        "title": "System Architecture RFC",
        "source_type": "design_doc",
        "access_policy": "internal",
    }

    res = await async_client.post(f"/api/v1/review-cycles/{cycle_id}/evidence/upload", files=files, data=data)
    assert res.status_code == 201
    ev = res.json()
    assert ev["title"] == "System Architecture RFC"
    assert ev["source_type"] == "design_doc"
    assert ev["storage_reference"] is not None
    assert os.path.exists(ev["storage_reference"])
    assert ev["metadata_json"]["detected_format"] == "pdf"


@pytest.mark.asyncio
async def test_upload_docx_evidence(async_client: AsyncClient, setup_upload_env):
    env = setup_upload_env
    cycle_id = env["cycle_id"]

    docx_bytes = generate_docx_bytes("Security Audit Findings", "All endpoints verified and hardened.")
    files = {
        "file": (
            "audit.docx",
            docx_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    data = {
        "cycle_id": cycle_id,
        "title": "Security Audit Report",
        "source_type": "document",
    }

    res = await async_client.post("/api/v1/evidence/upload", files=files, data=data)
    assert res.status_code == 201
    ev = res.json()
    assert ev["metadata_json"]["detected_format"] == "docx"
    assert "Security Audit Findings" in ev["metadata_json"]["snippet"]


@pytest.mark.asyncio
async def test_upload_csv_evidence_sanitizes_formulas(async_client: AsyncClient, setup_upload_env):
    env = setup_upload_env
    cycle_id = env["cycle_id"]

    csv_content = b"metric,value\nlatency,=cmd|'/C calc'!A0\nthroughput,15000\n"
    files = {"file": ("metrics.csv", csv_content, "text/csv")}
    data = {
        "cycle_id": cycle_id,
        "title": "Benchmark Metrics",
        "source_type": "metric",
    }

    res = await async_client.post("/api/v1/evidence/upload", files=files, data=data)
    assert res.status_code == 201
    ev = res.json()
    assert ev["metadata_json"]["detected_format"] == "csv"
    assert ev["metadata_json"]["parser_metadata"]["formula_injection_sanitized"] is True
    assert "'=cmd" in ev["metadata_json"]["snippet"]


@pytest.mark.asyncio
async def test_upload_streaming_size_limit_exceeded(async_client: AsyncClient, setup_upload_env):
    env = setup_upload_env
    cycle_id = env["cycle_id"]

    # Temporarily reduce MAX_UPLOAD_SIZE_BYTES to 100 KB for test verification
    with patch.object(settings, "MAX_UPLOAD_SIZE_BYTES", 100 * 1024):
        oversized_data = b"x" * (120 * 1024)  # 120 KB > 100 KB limit
        files = {"file": ("large_file.txt", oversized_data, "text/plain")}
        data = {"cycle_id": cycle_id, "title": "Oversized upload"}

        res = await async_client.post("/api/v1/evidence/upload", files=files, data=data)
        assert res.status_code == 413
        assert "exceeds maximum upload size limit" in res.json()["detail"]


@pytest.mark.asyncio
async def test_upload_empty_file_rejected(async_client: AsyncClient, setup_upload_env):
    env = setup_upload_env
    cycle_id = env["cycle_id"]

    files = {"file": ("empty.txt", b"", "text/plain")}
    data = {"cycle_id": cycle_id, "title": "Empty File"}

    res = await async_client.post("/api/v1/evidence/upload", files=files, data=data)
    assert res.status_code == 422
    assert "empty" in res.json()["detail"].lower()


@pytest.mark.asyncio
async def test_upload_executable_file_rejected(async_client: AsyncClient, setup_upload_env):
    env = setup_upload_env
    cycle_id = env["cycle_id"]

    # Windows PE / DOS executable header
    exe_content = b"MZ\x90\x00\x03\x00\x00\x00\x04\x00\x00\x00\xff\xff\x00\x00"
    files = {"file": ("malicious.exe", exe_content, "application/x-msdownload")}
    data = {"cycle_id": cycle_id, "title": "Trojan Upload"}

    res = await async_client.post("/api/v1/evidence/upload", files=files, data=data)
    assert res.status_code == 422
    assert "Executable binaries are strictly prohibited" in res.json()["detail"]


@pytest.mark.asyncio
async def test_upload_path_traversal_isolated_safely(async_client: AsyncClient, setup_upload_env):
    env = setup_upload_env
    cycle_id = env["cycle_id"]

    file_content = b"Legitimate text content inside an exploit filename."
    # Client filename attempts directory traversal
    traversal_filename = "../../../../../../../tmp/pwned.txt"
    files = {"file": (traversal_filename, file_content, "text/plain")}
    data = {"cycle_id": cycle_id, "title": "Traversal Attempt"}

    res = await async_client.post("/api/v1/evidence/upload", files=files, data=data)
    assert res.status_code == 201
    ev = res.json()

    # The storage_reference MUST be inside the isolated uploads directory for this cycle
    storage_path = ev["storage_reference"]
    upload_root = os.path.abspath(settings.UPLOAD_DIR)
    assert storage_path.startswith(upload_root)
    assert str(cycle_id) in storage_path
    assert not storage_path.startswith("/tmp/pwned.txt")
    assert os.path.exists(storage_path)


@pytest.mark.asyncio
async def test_upload_mismatched_employee_id_rejected(async_client: AsyncClient, setup_upload_env):
    env = setup_upload_env
    cycle_id = env["cycle_id"]

    # Create another distinct employee
    other_emp_res = await async_client.post(
        "/api/v1/users",
        json={"email": f"other_{uuid.uuid4().hex[:8]}@example.com", "full_name": "Other Emp", "role": "employee"},
    )
    other_employee_id = other_emp_res.json()["id"]

    files = {"file": ("notes.txt", b"Notes", "text/plain")}
    data = {
        "cycle_id": cycle_id,
        "employee_id": other_employee_id,  # Mismatched employee ID!
        "title": "Mismatched Employee",
    }

    res = await async_client.post("/api/v1/evidence/upload", files=files, data=data)
    assert res.status_code == 422
    assert "does not match review cycle employee ID" in res.json()["detail"]
