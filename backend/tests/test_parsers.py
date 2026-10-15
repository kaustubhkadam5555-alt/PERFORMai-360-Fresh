import io

import docx
import pytest
from pypdf import PdfWriter

from app.parsers.base import ParsedDocument
from app.parsers.csv_parser import CSVParser
from app.parsers.docx_parser import DocxParser
from app.parsers.factory import detect_and_validate_format
from app.parsers.pdf_parser import PDFParser
from app.parsers.text_parser import TextParser


def create_sample_pdf(text_content: str = "PERFORMai-360 Evaluation Document") -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def create_sample_docx(title: str = "RFC 42 Architecture", body: str = "System design details.") -> bytes:
    doc = docx.Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(body)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Latency"
    table.cell(1, 1).text = "25ms"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestTextParser:
    def test_parse_utf8_text(self):
        content = "Line 1: Refactored database repository.\nLine 2: Added connection pooling."
        stream = io.BytesIO(content.encode("utf-8"))
        parser = TextParser()
        result = parser.parse(stream, filename="notes.txt")

        assert isinstance(result, ParsedDocument)
        assert "Refactored database repository" in result.text
        assert result.word_count == 10
        assert result.metadata["format"] == "txt"
        assert result.metadata["filename"] == "notes.txt"

    def test_parse_latin1_text(self):
        content = "Café au lait - Performance Review".encode("latin-1")
        stream = io.BytesIO(content)
        parser = TextParser()
        result = parser.parse(stream, filename="review.txt")

        assert "Café" in result.text
        assert result.word_count > 0

    def test_strips_null_and_control_bytes(self):
        content = "Clean text\x00with null\x07bytes\x08here."
        stream = io.BytesIO(content.encode("utf-8"))
        parser = TextParser()
        result = parser.parse(stream, filename="control.txt")

        assert "\x00" not in result.text
        assert "\x07" not in result.text
        assert "Clean text" in result.text

    def test_truncates_oversized_text(self):
        oversized = "A" * 70_000
        stream = io.BytesIO(oversized.encode("utf-8"))
        parser = TextParser()
        result = parser.parse(stream, filename="large.txt")

        assert len(result.text) <= 50_000
        assert result.metadata["is_truncated"] is True


class TestCsvParser:
    def test_parse_valid_csv(self):
        csv_data = "Task ID,Title,Status\nTASK-101,Implement parser,Completed\nTASK-102,Write tests,In Progress\n"
        stream = io.BytesIO(csv_data.encode("utf-8"))
        parser = CSVParser()
        result = parser.parse(stream, filename="tasks.csv")

        assert "TASK-101" in result.text
        assert result.metadata["row_count"] == 3
        assert result.metadata["col_count"] == 3

    def test_neutralizes_formula_injection(self):
        malicious_csv = (
            "Name,Score,Command\n"
            "Alice,95,\"=cmd|'/C calc'!A0\"\n"
            'Bob,88,"+SUM(1,2)"\n'
            "Charlie,75,-2+3\n"
            "Dave,60,\"@HYPERLINK('http://evil.com','click')\"\n"
            "Eve,50,\tcmd.exe\n"
            'Frank,40,"\rcalc"\n'
        )
        stream = io.BytesIO(malicious_csv.encode("utf-8"))
        parser = CSVParser()
        result = parser.parse(stream, filename="export.csv")

        # Dangerous characters should be safely prefixed with single quote
        assert "'=cmd|'/C calc'!A0" in result.text
        assert "'+SUM(1,2)" in result.text
        assert "'-2+3" in result.text
        assert "'@HYPERLINK" in result.text
        assert "'cmd.exe" in result.text or "'\tcmd.exe" in result.text
        assert "'calc" in result.text or "'\rcalc" in result.text

    def test_enforces_char_limit_on_large_csv(self):
        lines = ["id,val"] + [f"{i},{'x' * 200}" for i in range(300)]
        stream = io.BytesIO("\n".join(lines).encode("utf-8"))
        parser = CSVParser()
        result = parser.parse(stream, filename="large.csv")

        assert len(result.text) <= 55_000
        assert result.metadata["is_truncated"] is True


class TestPdfParser:
    def test_parse_valid_pdf(self):
        pdf_bytes = create_sample_pdf()
        stream = io.BytesIO(pdf_bytes)
        parser = PDFParser()
        result = parser.parse(stream, filename="sample.pdf")

        assert isinstance(result, ParsedDocument)
        assert result.page_count == 1
        assert result.metadata["format"] == "pdf"

    def test_rejects_corrupted_pdf(self):
        corrupted = b"%PDF-1.4\nCorrupted garbage data that cannot be parsed by PDF parser."
        stream = io.BytesIO(corrupted)
        parser = PDFParser()
        with pytest.raises(ValueError, match="Corrupted or invalid PDF file"):
            parser.parse(stream, filename="corrupt.pdf")


class TestDocxParser:
    def test_parse_valid_docx(self):
        docx_bytes = create_sample_docx(title="Sprint 14 Retrospective", body="All goals met.")
        stream = io.BytesIO(docx_bytes)
        parser = DocxParser()
        result = parser.parse(stream, filename="retro.docx")

        assert "Sprint 14 Retrospective" in result.text
        assert "All goals met" in result.text
        assert "Latency" in result.text
        assert result.metadata["paragraph_count"] > 0
        assert result.metadata["table_count"] == 1

    def test_rejects_corrupted_docx(self):
        corrupted = b"PK\x03\x04\x14\x00Corrupted ZIP header that is not a valid DOCX"
        stream = io.BytesIO(corrupted)
        parser = DocxParser()
        with pytest.raises(ValueError, match="Corrupted or invalid DOCX file"):
            parser.parse(stream, filename="corrupt.docx")


class TestParserFactoryAndValidation:
    def test_detect_pdf(self):
        header = b"%PDF-1.7\n%\xe2\xe3\xcf\xd3"
        fmt, parser = detect_and_validate_format("doc.pdf", "application/pdf", header)
        assert fmt == "pdf"
        assert isinstance(parser, PDFParser)

    def test_detect_docx(self):
        header = b"PK\x03\x04\x14\x00\x06\x00"
        fmt, parser = detect_and_validate_format(
            "doc.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            header,
        )
        assert fmt == "docx"
        assert isinstance(parser, DocxParser)

    def test_detect_csv(self):
        header = b"col1,col2,col3\nval1,val2,val3\n"
        fmt, parser = detect_and_validate_format("metrics.csv", "text/csv", header)
        assert fmt == "csv"
        assert isinstance(parser, CSVParser)

    def test_detect_text(self):
        header = b"# Project README\nDetailed project notes."
        fmt, parser = detect_and_validate_format("notes.txt", "text/plain", header)
        assert fmt == "txt"
        assert isinstance(parser, TextParser)

    def test_rejects_executable_dos_mz(self):
        header = b"MZ\x90\x00\x03\x00\x00\x00\x04\x00"
        with pytest.raises(ValueError, match="Executable binaries are strictly prohibited"):
            detect_and_validate_format("evil.exe", "application/octet-stream", header)

    def test_rejects_executable_elf(self):
        header = b"\x7fELF\x02\x01\x01\x00"
        with pytest.raises(ValueError, match="Executable binaries are strictly prohibited"):
            detect_and_validate_format("evil.elf", "application/octet-stream", header)

    def test_rejects_executable_macho(self):
        header = b"\xcf\xfa\xed\xfe\x07\x00\x00\x01"
        with pytest.raises(ValueError, match="Executable binaries are strictly prohibited"):
            detect_and_validate_format("evil.dylib", "application/octet-stream", header)

    def test_rejects_script_shebang(self):
        header = b"#!/bin/bash\nrm -rf /\n"
        with pytest.raises(ValueError, match="Script files with shebang headers are strictly prohibited"):
            detect_and_validate_format("script.sh", "text/x-shellscript", header)

    def test_rejects_unsupported_format(self):
        header = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR"
        with pytest.raises(ValueError, match="Unsupported file format"):
            detect_and_validate_format("image.png", "image/png", header)
